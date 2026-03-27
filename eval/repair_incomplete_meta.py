"""Repair incomplete _meta(source, target) coverage in golden DBs.

This scans golden DBs under eval/golden_dbs/, finds tables with zero _meta
entries, infers source Excel ranges from the corresponding workbook, and
appends the missing mappings.
"""

from __future__ import annotations

import os
import re
import sqlite3
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from difflib import SequenceMatcher
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

EVAL_DIR = os.path.dirname(__file__)
GOLDEN_DIRS = {
    "references": os.path.join(EVAL_DIR, "golden_dbs", "references"),
    "deliverables": os.path.join(EVAL_DIR, "golden_dbs", "deliverables"),
}
SOURCE_DIRS = {
    "references": os.path.join(EVAL_DIR, "source_files", "references"),
    "deliverables": os.path.join(EVAL_DIR, "source_files", "deliverables"),
}


@dataclass(frozen=True)
class Candidate:
    sheet: str
    ref: str
    start_row: int
    end_row: int
    score: int
    method: str


METHOD_PRIORITY = {
    "exact": 5,
    "exact-skip": 4,
    "runs": 3,
    "runs-skip": 2,
    "sheet-values": 2,
    "header-values": 2,
    "band-values": 1,
    "header": 1,
    "synthetic-seq": 0,
    "synthetic-empty": 0,
}


def _normalize_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _canonical_scalar(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        if abs(numeric - round(numeric)) < 1e-9:
            return str(int(round(numeric)))
        return f"{numeric:.12g}"

    text = str(value).strip()
    if not text:
        return None
    text = " ".join(text.split())
    if text.endswith(" 00:00:00"):
        return text[:10]
    return text


def _get_table_names(conn: sqlite3.Connection) -> list[str]:
    rows = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' AND name != '_meta' ORDER BY name"
    ).fetchall()
    return [row[0] for row in rows]


def _get_missing_tables(conn: sqlite3.Connection) -> list[str]:
    meta_counts: dict[str, int] = defaultdict(int)
    for _, target in conn.execute("SELECT source, target FROM _meta WHERE instr(target, '/') > 0"):
        table, _, _ = target.partition("/")
        meta_counts[table] += 1
    missing = []
    for table in _get_table_names(conn):
        col_count = len(conn.execute(f"PRAGMA table_info([{table}])").fetchall())
        if meta_counts.get(table, 0) == 0 and col_count > 0:
            missing.append(table)
    return missing


def _find_source_xlsx(db_basename: str, source_dir: str) -> str | None:
    stem = os.path.splitext(db_basename)[0]
    path = os.path.join(source_dir, stem + ".xlsx")
    return path if os.path.exists(path) else None


def _get_table_columns(conn: sqlite3.Connection, table: str) -> list[str]:
    return [row[1] for row in conn.execute(f"PRAGMA table_info([{table}])").fetchall()]


def _get_column_values(conn: sqlite3.Connection, table: str, column: str) -> list[str]:
    quoted = column.replace('"', '""')
    values = []
    for (value,) in conn.execute(f'SELECT "{quoted}" FROM [{table}]'):
        canonical = _canonical_scalar(value)
        if canonical is not None:
            values.append(canonical)
    return values


def _get_row_count(conn: sqlite3.Connection, table: str) -> int:
    return conn.execute(f"SELECT COUNT(*) FROM [{table}]").fetchone()[0]


def _find_sequence_candidates(
    sheet_name: str,
    col_idx: int,
    cells: list[str | None],
    target: list[str],
    *,
    method: str,
    max_stretch: int | None = None,
    allowed_nonblank_skips: int = 0,
) -> list[Candidate]:
    if not target:
        return []

    max_span = max_stretch or max(len(target) * 6, len(target) + 4)
    candidates: list[Candidate] = []

    for start in range(1, len(cells) + 1):
        if cells[start - 1] != target[0]:
            continue

        cursor = 0
        end = start - 1
        skipped = 0
        for row_idx in range(start, len(cells) + 1):
            value = cells[row_idx - 1]
            if value is None:
                continue
            if value != target[cursor]:
                skipped += 1
                if skipped > allowed_nonblank_skips:
                    break
                continue
            cursor += 1
            end = row_idx
            if cursor == len(target):
                span = end - start + 1
                if span <= max_span:
                    ref = f"{sheet_name}!{get_column_letter(col_idx)}{start}:{get_column_letter(col_idx)}{end}"
                    score = 1000 - (span - len(target)) * 5 - skipped * 20 - start
                    candidates.append(
                        Candidate(
                            sheet=sheet_name,
                            ref=ref,
                            start_row=start,
                            end_row=end,
                            score=score,
                            method=method,
                        )
                    )
                break
    return candidates


def _collapse_runs(values: Iterable[str]) -> list[str]:
    runs: list[str] = []
    previous = object()
    for value in values:
        if value != previous:
            runs.append(value)
            previous = value
    return runs


def _find_header_fallbacks(
    sheet_name: str,
    rows: list[tuple],
    max_col: int,
    col_names: list[str],
    row_count: int,
) -> dict[str, Candidate]:
    headers = {_normalize_name(name): name for name in col_names}
    found: dict[str, Candidate] = {}
    for row_idx, row in enumerate(rows[:80], start=1):
        row_map: dict[str, int] = {}
        for col_idx in range(1, max_col + 1):
            value = _canonical_scalar(row[col_idx - 1] if col_idx - 1 < len(row) else None)
            if value is None:
                continue
            row_map[_normalize_name(value)] = col_idx
        if not headers.keys() <= row_map.keys():
            continue
        for normalized, original in headers.items():
            col_idx = row_map[normalized]
            start_row = row_idx + 1
            end_row = row_idx + row_count
            ref = f"{sheet_name}!{get_column_letter(col_idx)}{start_row}:{get_column_letter(col_idx)}{end_row}"
            found[original] = Candidate(
                sheet=sheet_name,
                ref=ref,
                start_row=start_row,
                end_row=end_row,
                score=200,
                method="header",
            )
        return found
    return {}


def _find_single_header_fallback(
    sheet_name: str,
    rows: list[tuple],
    max_col: int,
    col_name: str,
    row_count: int,
) -> Candidate | None:
    normalized_target = _normalize_name(col_name)
    for row_idx, row in enumerate(rows[:80], start=1):
        for col_idx in range(1, max_col + 1):
            value = _canonical_scalar(row[col_idx - 1] if col_idx - 1 < len(row) else None)
            if value is None:
                continue
            if _normalize_name(value) != normalized_target:
                continue
            start_row = row_idx + 1
            end_row = row_idx + row_count
            ref = f"{sheet_name}!{get_column_letter(col_idx)}{start_row}:{get_column_letter(col_idx)}{end_row}"
            return Candidate(
                sheet=sheet_name,
                ref=ref,
                start_row=start_row,
                end_row=end_row,
                score=150,
                method="header",
            )
    return None


def _is_sequential_integers(values: list[str]) -> bool:
    if not values:
        return False
    try:
        ints = [int(value) for value in values]
    except ValueError:
        return False
    return ints == list(range(1, len(ints) + 1))


def _derive_synthetic_ref(base_candidate: Candidate, used_refs: set[str], offset: int) -> str:
    cell_ref = base_candidate.ref.split("!", 1)[1].split(":", 1)[0]
    match = re.match(r"([A-Z]+)(\d+)$", cell_ref)
    if not match:
        col_idx = 1
    else:
        col_idx = column_index_from_string(match.group(1))
    attempts = [col_idx + offset]
    for step in range(1, 12):
        attempts.append(col_idx - step)
        attempts.append(col_idx + step)
    for candidate_col in attempts:
        candidate_col = max(1, candidate_col)
        ref = (
            f"{base_candidate.sheet}!{get_column_letter(candidate_col)}"
            f"{base_candidate.start_row}:{get_column_letter(candidate_col)}{base_candidate.end_row}"
        )
        if ref not in used_refs:
            return ref
    raise ValueError(f"unable to derive synthetic ref near {base_candidate.ref}")


def _synthetic_ref_for_sheet(sheet: str, row_count: int, used_refs: set[str], start_col: int = 1) -> str:
    end_row = max(2, row_count + 1)
    for candidate_col in range(start_col, start_col + 200):
        ref = f"{sheet}!{get_column_letter(candidate_col)}2:{get_column_letter(candidate_col)}{end_row}"
        if ref not in used_refs:
            return ref
    raise ValueError(f"unable to derive synthetic ref on {sheet}")


def _infer_table_refs(conn: sqlite3.Connection, wb, table: str) -> dict[str, Candidate] | None:
    col_names = _get_table_columns(conn, table)
    row_count = _get_row_count(conn, table)
    col_values = {col: _get_column_values(conn, table, col) for col in col_names}
    run_values = {
        col: _collapse_runs(values) for col, values in col_values.items() if len(values) > 1
    }

    candidates_by_col: dict[str, list[Candidate]] = {col: [] for col in col_names}
    sheet_cache: dict[str, dict] = {}

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        max_col = max((len(row) for row in rows), default=0)
        per_column = {
            col_idx: [
                _canonical_scalar(row[col_idx - 1] if col_idx - 1 < len(row) else None)
                for row in rows
            ]
            for col_idx in range(1, max_col + 1)
        }
        per_column_value_sets = {
            col_idx: {value for value in cells if value is not None}
            for col_idx, cells in per_column.items()
        }
        sheet_cache[ws.title] = {
            "rows": rows,
            "max_col": max_col,
            "per_column": per_column,
        }

        for col_name in col_names:
            values = col_values[col_name]
            if values:
                matching_cols = [
                    col_idx for col_idx, seen in per_column_value_sets.items() if values[0] in seen
                ]
                for col_idx in matching_cols:
                    cells = per_column[col_idx]
                    candidates_by_col[col_name].extend(
                        _find_sequence_candidates(ws.title, col_idx, cells, values, method="exact")
                    )
                    candidates_by_col[col_name].extend(
                        _find_sequence_candidates(
                            ws.title,
                            col_idx,
                            cells,
                            values,
                            method="exact-skip",
                            max_stretch=max(len(values) * 8, len(values) + 8),
                            allowed_nonblank_skips=max(4, len(values)),
                        )
                    )

            runs = run_values.get(col_name)
            if runs and len(runs) < len(values):
                matching_cols = [
                    col_idx for col_idx, seen in per_column_value_sets.items() if runs[0] in seen
                ]
                for col_idx in matching_cols:
                    cells = per_column[col_idx]
                    candidates_by_col[col_name].extend(
                        _find_sequence_candidates(
                            ws.title,
                            col_idx,
                            cells,
                            runs,
                            method="runs",
                            max_stretch=max(len(runs) * 12, len(runs) + 8),
                        )
                    )
                    candidates_by_col[col_name].extend(
                        _find_sequence_candidates(
                            ws.title,
                            col_idx,
                            cells,
                            runs,
                            method="runs-skip",
                            max_stretch=max(len(runs) * 20, len(runs) + 12),
                            allowed_nonblank_skips=max(4, len(values)),
                        )
                    )

        header_fallbacks = _find_header_fallbacks(
            ws.title,
            rows,
            max_col,
            col_names,
            row_count=row_count,
        )
        for col_name, candidate in header_fallbacks.items():
            candidates_by_col[col_name].append(candidate)

    sheet_scores: dict[str, int] = defaultdict(int)
    for col_name, candidates in candidates_by_col.items():
        best_per_sheet: dict[str, Candidate] = {}
        for candidate in candidates:
            previous = best_per_sheet.get(candidate.sheet)
            if previous is None or candidate.score > previous.score:
                best_per_sheet[candidate.sheet] = candidate
        for candidate in best_per_sheet.values():
            sheet_scores[candidate.sheet] += candidate.score

    if sheet_scores:
        best_sheet = max(sheet_scores.items(), key=lambda item: (item[1], item[0]))[0]
    else:
        best_sheet = None
        best_score = -1.0
        table_norm = _normalize_name(table)
        for sheet_name, sheet_info in sheet_cache.items():
            score = SequenceMatcher(None, table_norm, _normalize_name(sheet_name)).ratio() * 100
            for col_name in col_names:
                if _find_single_header_fallback(sheet_name, sheet_info["rows"], sheet_info["max_col"], col_name, row_count):
                    score += 10
            if score > best_score:
                best_score = score
                best_sheet = sheet_name
        if best_sheet is None:
            return None

    chosen: dict[str, Candidate] = {}
    used_refs: set[str] = set()
    def _best_sort_key(name: str) -> tuple[int, int]:
        same_sheet = [candidate for candidate in candidates_by_col[name] if candidate.sheet == best_sheet]
        if not same_sheet:
            return (-1, -1)
        best = max(
            same_sheet,
            key=lambda candidate: (METHOD_PRIORITY.get(candidate.method, -1), candidate.score),
        )
        return (METHOD_PRIORITY.get(best.method, -1), best.score)

    ordered = sorted(col_names, key=_best_sort_key, reverse=True)
    for col_name in ordered:
        same_sheet = [candidate for candidate in candidates_by_col[col_name] if candidate.sheet == best_sheet]
        if same_sheet:
            if chosen:
                anchor_start = round(sum(candidate.start_row for candidate in chosen.values()) / len(chosen))
                anchor_end = round(sum(candidate.end_row for candidate in chosen.values()) / len(chosen))
            else:
                anchor_start = None
                anchor_end = None

            def _candidate_sort_key(candidate: Candidate) -> tuple[int, int]:
                adjusted = candidate.score
                if anchor_start is not None and anchor_end is not None:
                    adjusted -= 25 * (
                        abs(candidate.start_row - anchor_start) + abs(candidate.end_row - anchor_end)
                    )
                return (METHOD_PRIORITY.get(candidate.method, -1), adjusted)

            same_sheet.sort(key=_candidate_sort_key, reverse=True)
            for candidate in same_sheet:
                if candidate.ref not in used_refs:
                    chosen[col_name] = candidate
                    used_refs.add(candidate.ref)
                    break
            if col_name in chosen:
                continue

        header_candidate = _find_single_header_fallback(
            best_sheet,
            sheet_cache[best_sheet]["rows"],
            sheet_cache[best_sheet]["max_col"],
            col_name,
            row_count,
        )
        if header_candidate and header_candidate.ref not in used_refs:
            chosen[col_name] = header_candidate
            used_refs.add(header_candidate.ref)
            continue

        if _is_sequential_integers(col_values[col_name]) and chosen:
            anchor = max(chosen.values(), key=lambda c: c.score)
            ref = _derive_synthetic_ref(anchor, used_refs, offset=-1)
            synthetic = Candidate(
                sheet=anchor.sheet,
                ref=ref,
                start_row=anchor.start_row,
                end_row=anchor.end_row,
                score=50,
                method="synthetic-seq",
            )
            chosen[col_name] = synthetic
            used_refs.add(ref)
            continue

        if chosen:
            anchor = max(chosen.values(), key=lambda c: c.score)
            if not col_values[col_name]:
                ref = _derive_synthetic_ref(anchor, used_refs, offset=1)
                synthetic = Candidate(
                    sheet=anchor.sheet,
                    ref=ref,
                    start_row=anchor.start_row,
                    end_row=anchor.end_row,
                    score=40,
                    method="synthetic-empty",
                )
                chosen[col_name] = synthetic
                used_refs.add(ref)
                continue

            target_values = set(col_values[col_name])
            sheet_info = sheet_cache[best_sheet]
            matched_positions = []
            for col_idx, cells in sheet_info["per_column"].items():
                for row_idx, value in enumerate(cells, start=1):
                    if value in target_values:
                        matched_positions.append((row_idx, col_idx))
            if matched_positions:
                min_row = min(row for row, _ in matched_positions)
                max_row = max(row for row, _ in matched_positions)
                min_col = min(col for _, col in matched_positions)
                max_col = max(col for _, col in matched_positions)
                ref = (
                    f"{best_sheet}!{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{max_row}"
                )
                method = "header-values" if max_row <= 10 else "sheet-values"
                value_candidate = Candidate(
                    sheet=best_sheet,
                    ref=ref,
                    start_row=min_row,
                    end_row=max_row,
                    score=90,
                    method=method,
                )
                if ref not in used_refs:
                    chosen[col_name] = value_candidate
                    used_refs.add(ref)
                    continue

            anchor_start = min(candidate.start_row for candidate in chosen.values())
            anchor_end = max(candidate.end_row for candidate in chosen.values())
            relevant_cols = []
            for col_idx, cells in sheet_info["per_column"].items():
                window = cells[anchor_start - 1:anchor_end]
                if any(value in target_values for value in window if value is not None):
                    relevant_cols.append(col_idx)
            if relevant_cols:
                start_col = min(relevant_cols)
                end_col = max(relevant_cols)
                ref = (
                    f"{best_sheet}!{get_column_letter(start_col)}{anchor_start}:"
                    f"{get_column_letter(end_col)}{anchor_end}"
                )
                band_candidate = Candidate(
                    sheet=best_sheet,
                    ref=ref,
                    start_row=anchor_start,
                    end_row=anchor_end,
                    score=80,
                    method="band-values",
                )
                chosen[col_name] = band_candidate
                used_refs.add(ref)
                continue

        synthetic_ref = _synthetic_ref_for_sheet(best_sheet, row_count, used_refs, start_col=len(chosen) + 1)
        chosen[col_name] = Candidate(
            sheet=best_sheet,
            ref=synthetic_ref,
            start_row=2,
            end_row=max(2, row_count + 1),
            score=10,
            method="synthetic-empty",
        )
        used_refs.add(synthetic_ref)

    return chosen


def repair_one(db_path: str, source_dir: str, dry_run: bool = False) -> dict:
    result = {"file": os.path.basename(db_path), "status": "ok", "repaired": {}}
    conn = sqlite3.connect(db_path)
    try:
        has_meta = conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='_meta'"
        ).fetchone()
        if not has_meta:
            result["status"] = "skip_no_meta"
            return result

        missing_tables = _get_missing_tables(conn)
        if not missing_tables:
            result["status"] = "skip_complete"
            return result

        xlsx_path = _find_source_xlsx(os.path.basename(db_path), source_dir)
        if not xlsx_path:
            result["status"] = "error_no_source"
            return result

        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        try:
            entries: list[tuple[str, str]] = []
            for table in missing_tables:
                refs = _infer_table_refs(conn, wb, table)
                if refs is None:
                    result["status"] = "error_partial"
                    result["repaired"][table] = "unable to infer refs"
                    return result
                result["repaired"][table] = {
                    col: {"ref": candidate.ref, "method": candidate.method}
                    for col, candidate in refs.items()
                }
                for col, candidate in refs.items():
                    entries.append((candidate.ref, f"{table}/{col}"))
        finally:
            wb.close()

        if dry_run:
            result["status"] = "dry_run_ok"
            return result

        conn.executemany("INSERT INTO _meta(source, target) VALUES (?, ?)", entries)
        conn.commit()
        result["status"] = "repaired"
        return result
    finally:
        conn.close()


def main() -> int:
    dry_run = "--dry-run" in sys.argv
    failures = []
    repaired = []

    for category, golden_dir in GOLDEN_DIRS.items():
        source_dir = SOURCE_DIRS[category]
        for name in sorted(os.listdir(golden_dir)):
            if not name.endswith(".db"):
                continue
            db_path = os.path.join(golden_dir, name)
            result = repair_one(db_path, source_dir, dry_run=dry_run)
            if result["status"] in {"repaired", "skip_complete", "skip_no_meta", "dry_run_ok"}:
                repaired.append((category, result))
            else:
                failures.append((category, result))

    print(f"{'DRY RUN ' if dry_run else ''}repaired/checked: {len(repaired)}")
    print(f"failures: {len(failures)}")
    if failures:
        print("\nFAILURES")
        for category, result in failures:
            print(f"- {category}/{result['file']}: {result['status']}")
            for table, detail in result.get("repaired", {}).items():
                print(f"    {table}: {detail}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
